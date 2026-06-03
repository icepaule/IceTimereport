"""
Generate the 'office' Excel report with ArbZG-corrected hours.
No project details, no violations visible. Clean for the boss.
"""
import os
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from azg import CorrectedDay
from holidays import is_holiday, is_weekend

GRAY_FILL = PatternFill(start_color="e0e0e0", end_color="e0e0e0", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
BLUE_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

DAY_NAMES_DE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
MONTH_NAMES_DE = [
    "", "Januar", "Februar", "März", "April", "Mai", "Juni",
    "Juli", "August", "September", "Oktober", "November", "Dezember",
]

HOURS_PER_WEEK = float(os.environ.get("HOURS_PER_WEEK", "39"))
HOURS_PER_DAY = HOURS_PER_WEEK / 5
VACATION_DAYS = int(os.environ.get("VACATION_DAYS", "30"))


def generate(
    year: int,
    corrected_days: list[CorrectedDay],
    output_dir: str,
    employee_name: str = "",
    employee_role: str = "",
    state: str = "BY",
    prior_overtime: float = 0.0,
    vacation_carryover: int = 0,
) -> str:
    """Generate office-safe Excel report. Returns output file path."""
    wb = Workbook()
    wb.remove(wb.active)

    day_lookup = {d.date: d for d in corrected_days}
    today = date.today()

    yearly_actual = 0.0
    yearly_target = 0.0
    yearly_vacation_used = 0
    yearly_sick_days = 0

    for month_num in range(1, 13):
        ws = wb.create_sheet(title=MONTH_NAMES_DE[month_num])

        # Header
        ws.merge_cells("A1:H1")
        ws["A1"] = f"Arbeitszeitnachweis {MONTH_NAMES_DE[month_num]} {year}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Mitarbeiter: {employee_name}"
        ws["A2"].font = BOLD
        ws["A3"] = f"Funktion: {employee_role}"
        ws["A3"].font = BOLD
        ws["A4"] = f"Soll: {HOURS_PER_WEEK}h/Woche"
        ws["A4"].font = BOLD

        headers = ["Datum", "Tag", "Typ", "Beginn", "Ende", "Pause (min)", "Ist (h)", "Soll (h)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        row = 7
        month_actual = 0.0
        month_target = 0.0
        month_vacation = 0
        month_sick = 0

        first_day = date(year, month_num, 1)
        if month_num == 12:
            last_day = date(year, 12, 31)
        else:
            last_day = date(year, month_num + 1, 1) - timedelta(days=1)

        current = first_day
        while current <= last_day:
            cd = day_lookup.get(current)
            weekend = is_weekend(current)
            holiday = is_holiday(current, state)

            if cd:
                day_type = cd.day_type
                hours = cd.corrected_hours
                start = cd.start_time
                end = cd.end_time
                pause = cd.pause_minutes if cd.pause_minutes else ""
            else:
                day_type = "Feiertag" if holiday else ("Samstag" if current.weekday() == 5 else "Sonntag" if current.weekday() == 6 else "Arbeit")
                hours = 0
                start = ""
                end = ""
                pause = ""

            # Future days: no target, no counters
            is_future = current > today

            # Target hours
            if is_future:
                target = 0
            elif day_type in ("Urlaub", "Sonderurlaub", "Krank", "Gleittag"):
                target = HOURS_PER_DAY
            elif weekend or holiday:
                target = 0
            else:
                target = HOURS_PER_DAY

            if not is_future:
                if day_type == "Urlaub":
                    month_vacation += 1
                elif day_type == "Krank":
                    month_sick += 1

            ws.cell(row=row, column=1, value=current.strftime("%d.%m.%Y")).border = THIN_BORDER
            ws.cell(row=row, column=2, value=DAY_NAMES_DE[current.weekday()]).border = THIN_BORDER
            ws.cell(row=row, column=3, value=day_type).border = THIN_BORDER
            ws.cell(row=row, column=4, value=start).border = THIN_BORDER
            ws.cell(row=row, column=5, value=end).border = THIN_BORDER
            ws.cell(row=row, column=6, value=pause).border = THIN_BORDER
            cell_h = ws.cell(row=row, column=7, value=round(hours, 2) if hours else "")
            cell_h.border = THIN_BORDER
            cell_h.number_format = "0.00"
            cell_t = ws.cell(row=row, column=8, value=round(target, 2) if target else "")
            cell_t.border = THIN_BORDER
            cell_t.number_format = "0.00"

            # Formatting
            if weekend or holiday:
                for c in range(1, 9):
                    ws.cell(row=row, column=c).fill = GRAY_FILL
            elif day_type in ("Urlaub", "Sonderurlaub", "Krank", "Gleittag"):
                for c in range(1, 9):
                    ws.cell(row=row, column=c).fill = YELLOW_FILL

            month_actual += hours
            month_target += target
            row += 1
            current += timedelta(days=1)

        # Monthly summary
        row += 1
        ws.cell(row=row, column=6, value="Monats-Summe:").font = BOLD
        ws.cell(row=row, column=7, value=round(month_actual, 2)).font = BOLD
        ws.cell(row=row, column=7).number_format = "0.00"
        ws.cell(row=row, column=8, value=round(month_target, 2)).font = BOLD
        ws.cell(row=row, column=8).number_format = "0.00"

        diff = month_actual - month_target
        ws.cell(row=row + 1, column=6, value="Differenz:").font = BOLD
        cell_diff = ws.cell(row=row + 1, column=7, value=round(diff, 2))
        cell_diff.font = Font(bold=True, color="FF0000" if diff < 0 else "008000")
        cell_diff.number_format = "+0.00;-0.00;0.00"

        if month_vacation:
            ws.cell(row=row + 2, column=6, value="Urlaubstage:").font = BOLD
            ws.cell(row=row + 2, column=7, value=month_vacation).font = BOLD
        if month_sick:
            ws.cell(row=row + 3, column=6, value="Kranktage:").font = BOLD
            ws.cell(row=row + 3, column=7, value=month_sick).font = BOLD

        yearly_actual += month_actual
        yearly_target += month_target
        yearly_vacation_used += month_vacation
        yearly_sick_days += month_sick

        widths = [12, 5, 12, 8, 8, 12, 8, 8]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Summary sheet
    ws_sum = wb.create_sheet(title="Zusammenfassung", index=0)
    ws_sum["A1"] = f"Arbeitszeitnachweis {year} - Zusammenfassung"
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum["A2"] = f"Mitarbeiter: {employee_name}"
    ws_sum["A3"] = f"Funktion: {employee_role}"

    cutoff = min(date(year, 12, 31), today)
    ws_sum["A4"] = f"Stand: {cutoff.strftime('%d.%m.%Y')}"
    ws_sum["A4"].font = Font(italic=True, size=10)

    ws_sum["A5"] = "Gesamt Ist-Stunden:"
    ws_sum["B5"] = round(yearly_actual, 2)
    ws_sum["B5"].number_format = "0.00"
    ws_sum["A6"] = "Gesamt Soll-Stunden:"
    ws_sum["B6"] = round(yearly_target, 2)
    ws_sum["B6"].number_format = "0.00"
    year_overtime = round(yearly_actual - yearly_target, 2)
    ws_sum["A7"] = f"Überstunden {year}:"
    ws_sum["B7"] = year_overtime
    ws_sum["B7"].number_format = "+0.00;-0.00;0.00"
    ws_sum["B7"].font = Font(bold=True, size=12)

    if prior_overtime != 0:
        ws_sum["A8"] = "Übertrag Vorjahre:"
        ws_sum["B8"] = round(prior_overtime, 2)
        ws_sum["B8"].number_format = "+0.00;-0.00;0.00"

    ws_sum["A9"] = "Überstundenkonto gesamt:"
    ws_sum["A9"].font = BOLD
    ws_sum["B9"] = round(prior_overtime + year_overtime, 2)
    ws_sum["B9"].number_format = "+0.00;-0.00;0.00"
    ws_sum["B9"].font = Font(bold=True, size=14, color="FF0000" if (prior_overtime + year_overtime) < 0 else "008000")

    effective_vacation = yearly_vacation_used - vacation_carryover
    ws_sum["A11"] = "Urlaubskonto:"
    ws_sum["A11"].font = BOLD
    ws_sum["A12"] = "Anspruch:"
    ws_sum["B12"] = VACATION_DAYS
    ws_sum["A13"] = "Genommen:"
    ws_sum["B13"] = effective_vacation
    if vacation_carryover:
        ws_sum["C13"] = f"({vacation_carryover} Tage aus Vorjahr)"
        ws_sum["C13"].font = Font(italic=True, size=9)
    ws_sum["A14"] = "Resturlaub:"
    ws_sum["B14"] = VACATION_DAYS - effective_vacation
    ws_sum["B14"].font = Font(bold=True, size=12)

    ws_sum["A16"] = "Krankheitstage:"
    ws_sum["B16"] = yearly_sick_days

    ws_sum.column_dimensions["A"].width = 25
    ws_sum.column_dimensions["B"].width = 15

    filepath = os.path.join(output_dir, f"Arbeitszeitnachweis_{year}.xlsx")
    os.makedirs(output_dir, exist_ok=True)
    wb.save(filepath)
    return filepath
