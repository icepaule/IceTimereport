"""
Generate the 'real' Excel report with actual hours and ArbZG violation column.
"""
import os
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from azg import DayInfo
from holidays import is_holiday, is_weekend

RED_FILL = PatternFill(start_color="FFcccc", end_color="FFcccc", fill_type="solid")
GREEN_FILL = PatternFill(start_color="ccFFcc", end_color="ccFFcc", fill_type="solid")
GRAY_FILL = PatternFill(start_color="e0e0e0", end_color="e0e0e0", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

DAY_NAMES_DE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
MONTH_NAMES_DE = [
    "", "Januar", "Februar", "März", "April", "Mai", "Juni",
    "Juli", "August", "September", "Oktober", "November", "Dezember",
]

HOURS_PER_WEEK = float(os.environ.get("HOURS_PER_WEEK", "39"))
HOURS_PER_DAY = HOURS_PER_WEEK / 5


def generate(
    year: int,
    days: list[DayInfo],
    output_dir: str,
    employee_name: str = "",
    employee_role: str = "",
    state: str = "BY",
    prior_overtime: float = 0.0,
    vacation_carryover: int = 0,
) -> str:
    """Generate real Excel report. Returns output file path."""
    wb = Workbook()
    wb.remove(wb.active)

    today = date.today()

    # Group days by month
    months: dict[int, list[DayInfo]] = {}
    for d in days:
        m = d.date.month
        months.setdefault(m, []).append(d)

    yearly_actual = 0.0
    yearly_target = 0.0
    yearly_violations: dict[str, int] = {}
    yearly_vacation_used = 0
    yearly_sick_days = 0

    for month_num in range(1, 13):
        month_days = months.get(month_num, [])
        ws = wb.create_sheet(title=MONTH_NAMES_DE[month_num])

        # Header
        ws.merge_cells("A1:H1")
        ws["A1"] = f"Arbeitszeitnachweis {MONTH_NAMES_DE[month_num]} {year}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Mitarbeiter: {employee_name}"
        ws["A2"].font = BOLD
        ws["A3"] = f"Funktion: {employee_role}"
        ws["A3"].font = BOLD
        ws["A4"] = f"Soll: {HOURS_PER_WEEK}h/Woche ({HOURS_PER_DAY:.1f}h/Tag)"
        ws["A4"].font = BOLD

        # Column headers
        headers = ["Datum", "Tag", "Typ", "Projekt", "Beschreibung", "Ist (h)", "Soll (h)", "ArbZG"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        # Build lookup for days
        day_lookup = {d.date: d for d in month_days}

        row = 7
        month_actual = 0.0
        month_target = 0.0
        month_vacation = 0
        month_sick = 0

        # Iterate all calendar days
        first_day = date(year, month_num, 1)
        if month_num == 12:
            last_day = date(year, 12, 31)
        else:
            last_day = date(year, month_num + 1, 1) - timedelta(days=1)

        current = first_day
        while current <= last_day:
            day_info = day_lookup.get(current)
            weekend = is_weekend(current)
            holiday = is_holiday(current, state)

            actual = day_info.actual_hours if day_info else 0.0
            violations = day_info.violations if day_info else []

            # Determine type and target
            day_type = _get_day_type(day_info, weekend, holiday)
            is_future = current > today

            if is_future:
                target = 0
            elif day_type in ("Urlaub", "Sonderurlaub", "Krank", "Gleittag"):
                target = HOURS_PER_DAY
            elif weekend or holiday:
                target = 0
            else:
                target = HOURS_PER_DAY

            if not is_future:
                if day_type == "Urlaub":
                    month_vacation += 1
                elif day_type == "Krank":
                    month_sick += 1

            # Project & description (aggregate if multiple entries)
            project = ""
            description = ""
            if day_info and day_info.entries:
                projects = list(dict.fromkeys(e.project_name for e in day_info.entries if e.project_name))
                descs = list(dict.fromkeys(e.description for e in day_info.entries if e.description))
                project = ", ".join(projects[:3])
                description = ", ".join(descs[:3])
                if len(projects) > 3:
                    project += " ..."
                if len(descs) > 3:
                    description += " ..."

            violation_str = ", ".join(violations) if violations else "OK" if actual > 0 else ""

            # Write row
            ws.cell(row=row, column=1, value=current.strftime("%d.%m.%Y")).border = THIN_BORDER
            ws.cell(row=row, column=2, value=DAY_NAMES_DE[current.weekday()]).border = THIN_BORDER
            ws.cell(row=row, column=3, value=day_type).border = THIN_BORDER
            ws.cell(row=row, column=4, value=project).border = THIN_BORDER
            ws.cell(row=row, column=5, value=description).border = THIN_BORDER
            cell_actual = ws.cell(row=row, column=6, value=round(actual, 2) if actual else "")
            cell_actual.border = THIN_BORDER
            cell_actual.number_format = "0.00"
            cell_target = ws.cell(row=row, column=7, value=round(target, 2) if target else "")
            cell_target.border = THIN_BORDER
            cell_target.number_format = "0.00"
            cell_azg = ws.cell(row=row, column=8, value=violation_str)
            cell_azg.border = THIN_BORDER

            # Formatting
            if violations:
                for col in range(1, 9):
                    ws.cell(row=row, column=col).fill = RED_FILL
                # Track violations
                for v in violations:
                    key = v.split("(")[0].strip()
                    yearly_violations[key] = yearly_violations.get(key, 0) + 1
            elif weekend or holiday:
                for col in range(1, 9):
                    ws.cell(row=row, column=col).fill = GRAY_FILL
            elif day_type in ("Urlaub", "Sonderurlaub", "Krank", "Gleittag"):
                for col in range(1, 9):
                    ws.cell(row=row, column=col).fill = YELLOW_FILL

            month_actual += actual
            month_target += target
            row += 1
            current += timedelta(days=1)

        # Monthly summary
        row += 1
        ws.cell(row=row, column=5, value="Monats-Summe:").font = BOLD
        ws.cell(row=row, column=6, value=round(month_actual, 2)).font = BOLD
        ws.cell(row=row, column=6).number_format = "0.00"
        ws.cell(row=row, column=7, value=round(month_target, 2)).font = BOLD
        ws.cell(row=row, column=7).number_format = "0.00"
        diff = month_actual - month_target
        ws.cell(row=row + 1, column=5, value="Differenz:").font = BOLD
        cell_diff = ws.cell(row=row + 1, column=6, value=round(diff, 2))
        cell_diff.font = Font(bold=True, color="FF0000" if diff < 0 else "008000")
        cell_diff.number_format = "+0.00;-0.00;0.00"

        yearly_actual += month_actual
        yearly_target += month_target
        yearly_vacation_used += month_vacation
        yearly_sick_days += month_sick

        # Column widths
        widths = [12, 5, 10, 25, 35, 8, 8, 25]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Summary sheet
    ws_sum = wb.create_sheet(title="Zusammenfassung", index=0)
    ws_sum["A1"] = f"Arbeitszeitnachweis {year} - Zusammenfassung"
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum["A2"] = f"Mitarbeiter: {employee_name}"
    ws_sum["A3"] = f"Funktion: {employee_role}"

    cutoff = min(date(year, 12, 31), today)
    ws_sum["A4"] = f"Stand: {cutoff.strftime('%d.%m.%Y')}"
    ws_sum["A4"].font = Font(italic=True, size=10)

    ws_sum["A5"] = "Gesamt Ist-Stunden:"
    ws_sum["B5"] = round(yearly_actual, 2)
    ws_sum["B5"].number_format = "0.00"
    ws_sum["A6"] = "Gesamt Soll-Stunden:"
    ws_sum["B6"] = round(yearly_target, 2)
    ws_sum["B6"].number_format = "0.00"
    year_overtime = round(yearly_actual - yearly_target, 2)
    ws_sum["A7"] = f"Überstunden {year}:"
    ws_sum["B7"] = year_overtime
    ws_sum["B7"].number_format = "+0.00;-0.00;0.00"
    ws_sum["B7"].font = Font(bold=True, size=12)

    if prior_overtime != 0:
        ws_sum["A8"] = "Übertrag Vorjahre:"
        ws_sum["B8"] = round(prior_overtime, 2)
        ws_sum["B8"].number_format = "+0.00;-0.00;0.00"

    ws_sum["A9"] = "Überstundenkonto gesamt:"
    ws_sum["A9"].font = BOLD
    ws_sum["B9"] = round(prior_overtime + year_overtime, 2)
    ws_sum["B9"].number_format = "+0.00;-0.00;0.00"
    ws_sum["B9"].font = Font(bold=True, size=14, color="FF0000" if (prior_overtime + year_overtime) < 0 else "008000")

    VACATION_DAYS = int(os.environ.get("VACATION_DAYS", "30"))
    effective_vacation = yearly_vacation_used - vacation_carryover
    ws_sum["A11"] = "Urlaubskonto:"
    ws_sum["A11"].font = BOLD
    ws_sum["A12"] = "Anspruch:"
    ws_sum["B12"] = VACATION_DAYS
    ws_sum["A13"] = "Genommen:"
    ws_sum["B13"] = effective_vacation
    if vacation_carryover:
        ws_sum["C13"] = f"({vacation_carryover} Tage aus Vorjahr)"
        ws_sum["C13"].font = Font(italic=True, size=9)
    ws_sum["A14"] = "Resturlaub:"
    ws_sum["B14"] = VACATION_DAYS - effective_vacation
    ws_sum["B14"].font = Font(bold=True, size=12)

    ws_sum["A16"] = "Krankheitstage:"
    ws_sum["B16"] = yearly_sick_days

    if yearly_violations:
        ws_sum["A18"] = "ArbZG-Verstöße:"
        ws_sum["A18"].font = BOLD
        r = 19
        for violation_type, count in sorted(yearly_violations.items()):
            ws_sum[f"A{r}"] = violation_type
            ws_sum[f"B{r}"] = count
            ws_sum[f"A{r}"].fill = RED_FILL
            ws_sum[f"B{r}"].fill = RED_FILL
            r += 1
        ws_sum[f"A{r}"] = "Gesamt"
        ws_sum[f"A{r}"].font = BOLD
        ws_sum[f"B{r}"] = sum(yearly_violations.values())
        ws_sum[f"B{r}"].font = BOLD

    ws_sum.column_dimensions["A"].width = 25
    ws_sum.column_dimensions["B"].width = 15

    filepath = os.path.join(output_dir, f"Arbeitszeitnachweis_{year}_real.xlsx")
    os.makedirs(output_dir, exist_ok=True)
    wb.save(filepath)
    return filepath


def _get_day_type(day_info: DayInfo | None, weekend: bool, holiday: str | None) -> str:
    if holiday:
        return f"Feiertag"
    if weekend:
        return "Samstag" if day_info and day_info.date.weekday() == 5 else "Sonntag" if day_info and day_info.date.weekday() == 6 else "Wochenende"
    if day_info and day_info.entries:
        projects = {e.project_name.lower() for e in day_info.entries}
        for p in projects:
            if "sonderurlaub" in p:
                return "Sonderurlaub"
            if "urlaub" in p:
                return "Urlaub"
            if "krank" in p:
                return "Krank"
            if "gleittag" in p or "gleitzeit" in p:
                return "Gleittag"
    return "Arbeit"
